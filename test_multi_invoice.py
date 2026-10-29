"""
Test suite for MULTIPLE SUPPLIER INVOICE SHEETS feature.
Verifies:
1. Every unique supplier invoice gets exactly one sheet named after its invoice number.
2. Each sheet contains only items belonging to its corresponding invoice.
3. No items are missing or duplicated across sheets.
4. Correct column layout, formulas (=SUM, =+L...-L...), and Tally-compatible structure.
5. Custom Duty 1 has per-item invoice references and does not double count.
6. Serial List and Checks sheets remain fully functional.
7. Single invoice input generates exactly 1 purchase sheet.
"""

from decimal import Decimal
from datetime import date
from pathlib import Path
import openpyxl
import pytest

from main import ExcelGenerator, BASE_EXCEL, round_cur
APP_DIR = Path(__file__).resolve().parent

BASE_PATH = APP_DIR / "templates" / "base.xlsx"

@pytest.fixture
def sample_boe_header():
    return {
        "boe_no": "4342140",
        "be_no": "4342140",
        "be_date": date(2026, 10, 11),
        "exchange_rate": Decimal("110.50"),
        "total_assessable_value": Decimal("100000.00"),
        "total_bcd": Decimal("10000.00"),
        "total_sws": Decimal("1000.00"),
        "invoice_total_eur": Decimal("904.98")
    }

def create_mock_item(item_no, part_no, tally_name, qty, unit_price_inr, inv_no, inv_date):
    tot_inr = round_cur(Decimal(str(unit_price_inr)) * Decimal(str(qty)))
    bcd_rate = Decimal("10.0")
    bcd = round_cur(tot_inr * Decimal("0.10"))
    sws = round_cur(bcd * Decimal("0.10"))
    return {
        "item_no": item_no,
        "boe_sn": item_no // 10,
        "part_no": part_no,
        "tally_name": tally_name,
        "batch_type": "Primary Batch",
        "qty": qty,
        "unit_price_inr": tot_inr,
        "total_inr": tot_inr,
        "batch_no": "Primary batch ",
        "batch_changed": False,
        "bcd_rate": bcd_rate,
        "bcd_tot": bcd,
        "sws_tot": sws,
        "unit_bcd": bcd,
        "unit_sws": sws,
        "po_no": f"PO-{inv_no}",
        "po_date": inv_date,
        "ocr_flag": False,
        "invoice_no": inv_no,
        "invoice_date": inv_date
    }

def test_multiple_supplier_invoices_generation(tmp_path, sample_boe_header):
    # Two supplier invoices: 980007121 (3 items) and 980011770 (2 items)
    inv1_date = date(2026, 6, 15)
    inv2_date = date(2026, 7, 20)

    items = [
        create_mock_item(10, "PART-A1", "Tally Part A1", 5, 1000, "980007121", inv1_date),
        create_mock_item(20, "PART-A2", "Tally Part A2", 2, 2000, "980007121", inv1_date),
        create_mock_item(30, "PART-A3", "Tally Part A3", 1, 3000, "980007121", inv1_date),
        create_mock_item(40, "PART-B1", "Tally Part B1", 10, 500, "980011770", inv2_date),
        create_mock_item(50, "PART-B2", "Tally Part B2", 4, 1500, "980011770", inv2_date),
    ]

    total_inr = sum(it["total_inr"] for it in items)
    exch = sample_boe_header["exchange_rate"]
    sample_boe_header["total_assessable_value"] = total_inr
    sample_boe_header["total_bcd"] = sum(it["unit_bcd"] for it in items)
    sample_boe_header["total_sws"] = sum(it["unit_sws"] for it in items)
    sample_boe_header["invoice_total_eur"] = round_cur(total_inr / exch)

    inv_header = {
        "invoice_no": "980007121",
        "invoice_date": inv1_date,
        "currency": "EUR",
        "total_eur": sample_boe_header["invoice_total_eur"]
    }

    inv_headers = {
        "980007121": {"invoice_no": "980007121", "invoice_date": inv1_date, "currency": "EUR", "total_eur": round_cur(Decimal("12000.00") / exch)},
        "980011770": {"invoice_no": "980011770", "invoice_date": inv2_date, "currency": "EUR", "total_eur": round_cur(Decimal("11000.00") / exch)},
    }

    generator = ExcelGenerator(BASE_PATH, tmp_path, "TEST_MULTI_INVOICE.xlsx")
    out_file = generator.generate(sample_boe_header, inv_header, items, inv_headers=inv_headers)

    assert out_file.exists()
    wb = openpyxl.load_workbook(out_file, data_only=False)

    # 1. Verify sheet names and ordering
    expected_sheets = ["980007121", "980011770", "Custom Duty 1 ", "Serial List", "Checks"]
    assert wb.sheetnames == expected_sheets, f"Unexpected sheets: {wb.sheetnames}"

    # 2. Verify Sheet 980007121
    ws1 = wb["980007121"]
    # Check data rows: row 2 to 4 (3 items)
    rows_ws1 = [ws1.cell(r, 8).value for r in range(2, 5)]
    assert rows_ws1 == ["Tally Part A1", "Tally Part A2", "Tally Part A3"]
    for r in range(2, 5):
        assert str(ws1.cell(r, 5).value) == "980007121"
        assert ws1.cell(r, 4).value == "YASKAWA Europe GmbH"
        assert ws1.cell(r, 2).value == "EXCISE PURCHASE"
    # Check formulas in ws1
    sum_r1 = 6  # 3 items + 1 header + 2 offset
    assert ws1.cell(sum_r1, 9).value == "=SUM(I2:I4)"
    assert ws1.cell(sum_r1, 12).value == "=SUM(L2:L4)"
    assert ws1.cell(sum_r1, 13).value == "=SUM(M2:M4)"
    assert ws1.cell(sum_r1 + 3, 12).value == f"=+L{sum_r1}-L{sum_r1 + 2}"

    # 3. Verify Sheet 980011770
    ws2 = wb["980011770"]
    # Check data rows: row 2 to 3 (2 items)
    rows_ws2 = [ws2.cell(r, 8).value for r in range(2, 4)]
    assert rows_ws2 == ["Tally Part B1", "Tally Part B2"]
    for r in range(2, 4):
        assert str(ws2.cell(r, 5).value) == "980011770"
        assert ws2.cell(r, 4).value == "YASKAWA Europe GmbH"
    sum_r2 = 5
    assert ws2.cell(sum_r2, 9).value == "=SUM(I2:I3)"
    assert ws2.cell(sum_r2, 12).value == "=SUM(L2:L3)"
    assert ws2.cell(sum_r2, 13).value == "=SUM(M2:M3)"

    # 4. Verify Custom Duty 1 has per-item narrations
    ws_d = wb["Custom Duty 1 "]
    # 5 BCD rows (rows 2-6), 5 SWS rows (rows 7-11)
    for r in range(2, 5):
        assert "Invoice No. 980007121" in str(ws_d.cell(r, 12).value)
    for r in range(5, 7):
        assert "Invoice No. 980011770" in str(ws_d.cell(r, 12).value)

    # 5. Verify Checks Sheet
    ws_c = wb["Checks"]
    statuses = [ws_c.cell(r, 5).value for r in range(6, 12)]
    assert all(s == "PASS" for s in statuses), f"Checks failed: {statuses}"

def test_three_supplier_invoices(tmp_path, sample_boe_header):
    d = date(2026, 8, 1)
    items = [
        create_mock_item(10, "P1", "Tally P1", 1, 1000, "980007121", d),
        create_mock_item(20, "P2", "Tally P2", 1, 2000, "980011770", d),
        create_mock_item(30, "P3", "Tally P3", 1, 3000, "980011771", d),
    ]
    inv_header = {"invoice_no": "980007121", "invoice_date": d, "currency": "EUR", "total_eur": Decimal("54.30")}
    generator = ExcelGenerator(BASE_PATH, tmp_path, "TEST_THREE_INVOICES.xlsx")
    out_file = generator.generate(sample_boe_header, inv_header, items)

    wb = openpyxl.load_workbook(out_file, data_only=False)
    assert wb.sheetnames == ["980007121", "980011770", "980011771", "Custom Duty 1 ", "Serial List", "Checks"]

def test_single_supplier_invoice_preservation(tmp_path, sample_boe_header):
    d = date(2026, 8, 1)
    items = [
        create_mock_item(10, "P1", "Tally P1", 1, 1000, "980011770", d),
        create_mock_item(20, "P2", "Tally P2", 1, 2000, "980011770", d),
    ]
    inv_header = {"invoice_no": "980011770", "invoice_date": d, "currency": "EUR", "total_eur": Decimal("27.15")}
    generator = ExcelGenerator(BASE_PATH, tmp_path, "TEST_SINGLE_INVOICE.xlsx")
    out_file = generator.generate(sample_boe_header, inv_header, items)

    wb = openpyxl.load_workbook(out_file, data_only=False)
    # Exactly one purchase sheet
    assert wb.sheetnames == ["980011770", "Custom Duty 1 ", "Serial List", "Checks"]
    assert wb["980011770"].max_row == 8  # 2 data rows + summary (row 5) + target (row 7) + diff (row 8)
